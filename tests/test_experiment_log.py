"""Experiment ledger: field extraction, dataset catalog, xlsx append."""

from __future__ import annotations

import json

import pytest
from openpyxl import load_workbook

from conftest import submission_checks as all_pass_checks
from worldquant.api import SUBMISSION_CHECKS, SimulationStatus
from worldquant.config import AppConfig, StorageConfig
from worldquant.experiment_log import (
    LEDGER_COLUMNS,
    ExperimentLog,
    FieldCatalog,
    build_ledger,
    extract_field_names,
    failed_checks,
    write_ledger_from_results,
)
from worldquant.models import AlphaResult

SETTINGS = {
    "region": "USA", "universe": "TOP3000", "delay": 1, "decay": 8,
    "neutralization": "SUBINDUSTRY", "truncation": 0.02, "pasteurization": "ON",
    "nanHandling": "OFF", "unitHandling": "VERIFY", "instrumentType": "EQUITY",
    "language": "FASTEXPR",
}


def make_result(**overrides) -> AlphaResult:
    base = dict(
        alpha_id="a1",
        expression="-ts_rank(fn_liab_fair_val_l1_a, 126)",
        dedup_key="k1",
        status=SimulationStatus.COMPLETED,
        grade="AVERAGE",
        stage="IS",
        simulation_id="sim1",
        remote_alpha_id="kqVkNWmP",
        settings_json=json.dumps(SETTINGS),
        sharpe=1.47, fitness=1.20, turnover=0.0383, returns=0.0833,
        drawdown=0.0672, margin=0.00435, long_count=438, short_count=432,
        checks={
            "LOW_SHARPE": {"value": 1.47, "result": "PASS", "limit": 1.25},
            "LOW_FITNESS": {"value": 1.20, "result": "PASS", "limit": 1.0},
        },
        passed=True,
        created_at="2026-09-06T00:00:00+00:00",
        completed_at="2026-09-06T00:10:00+00:00",
    )
    base.update(overrides)
    return AlphaResult(**base)


def read_rows(path):
    workbook = load_workbook(path)
    sheet = workbook.active
    header = [cell.value for cell in sheet[1]]
    rows = []
    for raw in sheet.iter_rows(min_row=2, values_only=True):
        rows.append(dict(zip(header, raw)))
    return header, rows


class TestExtractFieldNames:
    def test_simple_expression(self):
        assert extract_field_names("rank(ts_delta(close, 5))") == ["close"]

    def test_fundamental_field(self):
        assert extract_field_names("-ts_rank(fn_liab_fair_val_l1_a, 126)") == [
            "fn_liab_fair_val_l1_a"
        ]

    def test_multiple_fields_keep_first_seen_order(self):
        assert extract_field_names("rank(close) - rank(volume) + rank(close)") == [
            "close", "volume",
        ]

    def test_function_names_are_not_fields(self):
        fields = extract_field_names(
            "ts_regression(returns, fn_assets_fair_val_l3_a, 60)"
        )
        assert "ts_regression" not in fields
        assert fields == ["returns", "fn_assets_fair_val_l3_a"]

    def test_string_literals_are_not_fields(self):
        fields = extract_field_names("group_neutralize(rank(returns), 'industry')")
        assert fields == ["returns"]
        assert "industry" not in fields

    def test_double_quoted_literals_are_not_fields(self):
        assert extract_field_names('group_neutralize(rank(returns), "subindustry")') == ["returns"]

    def test_named_arguments_are_not_fields(self):
        fields = extract_field_names("ts_regression(returns, cap, 60, RETTYPE=1)")
        assert "RETTYPE" not in fields
        assert fields == ["returns", "cap"]

    def test_comparison_operands_are_still_fields(self):
        fields = extract_field_names("if(close == open, volume, returns)")
        assert "close" in fields
        assert "open" in fields

    def test_numbers_are_not_fields(self):
        assert extract_field_names("rank(ts_delta(close, 126))") == ["close"]

    def test_empty_expression(self):
        assert extract_field_names("") == []

    def test_composite_score_field(self):
        assert extract_field_names("rank(multi_factor_static_score_derivative)") == [
            "multi_factor_static_score_derivative"
        ]


MULTI_STATEMENT = (
    "ey = ts_backfill(anl4_ebit_value, 40) / cap; "
    "raw = group_zscore(winsorize(ts_mean(ey - ts_mean(ey, 120), 20), std=3.0), industry); "
    "vol_state = ts_rank(ts_std_dev(returns, 20), 252); "
    "liq_state = ts_rank(volume, 20); "
    "trade_when((liq_state > 0.40) && (vol_state < 0.90), raw, vol_state > 0.97)"
)


class TestMultiStatementExpressions:
    """FASTEXPR lets an expression bind intermediates; those are not data fields."""

    def test_local_variables_are_excluded(self):
        fields = extract_field_names(MULTI_STATEMENT)
        for variable in ("ey", "raw", "vol_state", "liq_state"):
            assert variable not in fields

    def test_real_fields_are_still_found(self):
        assert extract_field_names(MULTI_STATEMENT) == [
            "anl4_ebit_value", "cap", "returns", "volume",
        ]

    def test_references_to_a_variable_are_excluded_too(self):
        # `x` is defined first, then referenced twice afterwards.
        assert extract_field_names("x = close; rank(x) + rank(x)") == ["close"]

    def test_variable_defined_after_its_first_use_still_excluded(self):
        # The assigned-name pass scans the whole expression first, so the order
        # of definition versus reference does not matter.
        assert extract_field_names("rank(tmp) + tmp; tmp = volume") == ["volume"]

    def test_named_arguments_are_not_treated_as_definitions(self):
        # std=3.0 sits inside winsorize(...), at depth >= 1, so it must not be
        # collected as a variable name and shadow a real field called "std".
        fields = extract_field_names("winsorize(rank(close), std=3.0)")
        assert fields == ["close"]
        assert "std" not in fields

    def test_comparison_operands_are_not_assignments(self):
        fields = extract_field_names(
            "trade_when((liq > 0.4) && (vol < 0.9), rank(close), vol > 0.97)"
        )
        assert "close" in fields
        assert "liq" in fields
        assert "vol" in fields

    def test_equality_comparison_keeps_both_operands(self):
        assert extract_field_names("if(close == open, volume, returns)") == [
            "close", "open", "volume", "returns",
        ]

    def test_string_literal_containing_a_paren_does_not_break_depth_tracking(self):
        fields = extract_field_names(
            "a = group_neutralize(rank(close), '('); rank(a) + rank(volume)"
        )
        assert fields == ["close", "volume"]

    def test_single_statement_expression_is_unaffected(self):
        assert extract_field_names("rank(ts_delta(close, 5))") == ["close"]


class TestFailedChecks:
    def test_only_non_passing_checks_are_reported(self):
        result = make_result(checks={
            "LOW_SHARPE": {"value": 1.92, "result": "PASS", "limit": 1.25},
            "LOW_FITNESS": {"value": 0.83, "result": "FAIL", "limit": 1.0},
            "HIGH_TURNOVER": {"value": 0.9466, "result": "FAIL", "limit": 0.7},
            "SELF_CORRELATION": {"value": None, "result": "PENDING", "limit": None},
        })
        failed = failed_checks(result)
        assert len(failed) == 2
        assert any(entry.startswith("LOW_FITNESS=0.83") for entry in failed)
        assert any("HIGH_TURNOVER" in entry for entry in failed)

    def test_all_passing_gives_an_empty_list(self):
        assert failed_checks(make_result()) == []

    def test_no_checks_at_all(self):
        assert failed_checks(make_result(checks={})) == []

    def test_malformed_check_entries_are_skipped(self):
        assert failed_checks(make_result(checks={"LOW_SHARPE": "not-a-dict"})) == []

    def test_the_resolved_checks_supersede_the_alpha_payload(self):
        # The alpha payload reports SELF_CORRELATION as PENDING forever. Reading
        # only that copy would leave `submittable=NO` in the ledger with no
        # reason attached, which is the one thing the column exists to prevent.
        result = make_result(
            checks={"SELF_CORRELATION": {"value": None, "result": "PENDING", "limit": 0.7}},
            submission_checks=all_pass_checks({"SELF_CORRELATION": "FAIL"}),
        )
        assert failed_checks(result) == ["SELF_CORRELATION:FAIL"]

    def test_a_pending_verdict_is_not_an_active_failure(self):
        # checks_passed reading 7/8 is what surfaces an unresolved check.
        result = make_result(submission_checks=all_pass_checks({"SELF_CORRELATION": "PENDING"}))
        assert failed_checks(result) == []


class TestFieldCatalog:
    def test_resolves_through_the_injected_fetcher(self, tmp_path):
        calls = []

        def fetch(field):
            calls.append(field)
            return {"id": field, "dataset": {"id": "model16", "name": "Fundamental Scores"}}

        catalog = FieldCatalog(tmp_path / "cat.json", fetch=fetch)
        assert catalog.dataset_for("multi_factor_static_score_derivative") == "model16"
        assert calls == ["multi_factor_static_score_derivative"]

    def test_second_lookup_is_served_from_memory(self, tmp_path):
        calls = []
        catalog = FieldCatalog(
            tmp_path / "cat.json",
            fetch=lambda f: calls.append(f) or {"dataset": {"id": "pv1"}},
        )
        catalog.dataset_for("close")
        catalog.dataset_for("close")
        assert calls == ["close"]

    def test_cache_survives_a_reopen(self, tmp_path):
        path = tmp_path / "cat.json"
        first = FieldCatalog(path, fetch=lambda f: {"dataset": {"id": "model16"}})
        first.dataset_for("some_field")

        calls = []
        second = FieldCatalog(path, fetch=lambda f: calls.append(f) or None)
        assert second.dataset_for("some_field") == "model16"
        assert calls == [], "a cached field must not hit the API again"

    def test_unresolvable_field_is_not_retried(self, tmp_path):
        calls = []
        catalog = FieldCatalog(tmp_path / "cat.json", fetch=lambda f: calls.append(f) or None)
        assert catalog.dataset_for("bogus") is None
        assert catalog.dataset_for("bogus") is None
        assert calls == ["bogus"]

    def test_fetcher_raising_is_tolerated(self, tmp_path):
        from worldquant.exceptions import APIError

        def boom(field):
            raise APIError("nope", status_code=500, url="/data-fields/x")

        catalog = FieldCatalog(tmp_path / "cat.json", fetch=boom)
        assert catalog.dataset_for("close") is None

    def test_no_fetcher_means_no_resolution(self, tmp_path):
        assert FieldCatalog(tmp_path / "cat.json").dataset_for("close") is None

    def test_datasets_for_dedupes_and_preserves_order(self, tmp_path):
        mapping = {"close": "pv1", "volume": "pv1", "cap": "fundamental6"}
        catalog = FieldCatalog(
            tmp_path / "cat.json", fetch=lambda f: {"dataset": {"id": mapping.get(f)}}
        )
        assert catalog.datasets_for(["close", "volume", "cap"]) == ["pv1", "fundamental6"]

    def test_dataset_as_a_plain_string_is_accepted(self, tmp_path):
        catalog = FieldCatalog(tmp_path / "cat.json", fetch=lambda f: {"dataset": "model16"})
        assert catalog.dataset_for("x") == "model16"

    def test_corrupt_cache_is_ignored_not_fatal(self, tmp_path):
        path = tmp_path / "cat.json"
        path.write_text("{not json", encoding="utf-8")
        catalog = FieldCatalog(path, fetch=lambda f: {"dataset": {"id": "pv1"}})
        assert catalog.dataset_for("close") == "pv1"


class TestExperimentLog:
    def test_creates_the_workbook_with_a_header(self, tmp_path):
        path = tmp_path / "exp.xlsx"
        ledger = ExperimentLog(path)
        assert ledger.record(make_result(), SETTINGS) is True

        header, rows = read_rows(path)
        assert header == list(LEDGER_COLUMNS)
        assert len(rows) == 1

    def test_appends_one_row_per_simulation(self, tmp_path):
        path = tmp_path / "exp.xlsx"
        ledger = ExperimentLog(path)
        for index in range(3):
            ledger.record(make_result(alpha_id=f"a{index}"), SETTINGS)

        _, rows = read_rows(path)
        assert [row["alpha_id"] for row in rows] == ["a0", "a1", "a2"]
        assert ledger.rows_written == 3

    def test_records_the_dataset_fields_and_parameters(self, tmp_path):
        path = tmp_path / "exp.xlsx"
        catalog = FieldCatalog(
            tmp_path / "cat.json",
            fetch=lambda f: {"dataset": {"id": "fundamental2", "name": "Report Footnotes"}},
        )
        ledger = ExperimentLog(path, catalog=catalog)
        ledger.record(make_result(), SETTINGS)

        _, rows = read_rows(path)
        row = rows[0]
        assert row["fields"] == "fn_liab_fair_val_l1_a"
        assert row["field_count"] == 1
        assert row["datasets"] == "fundamental2"
        assert row["decay"] == 8
        assert row["neutralization"] == "SUBINDUSTRY"
        assert row["truncation"] == 0.02
        assert row["region"] == "USA"
        assert row["nan_handling"] == "OFF"
        assert row["unit_handling"] == "VERIFY"

    def test_records_the_results_and_grade(self, tmp_path):
        path = tmp_path / "exp.xlsx"
        ExperimentLog(path).record(make_result(), SETTINGS)

        _, rows = read_rows(path)
        row = rows[0]
        assert row["grade"] == "AVERAGE"
        assert row["sharpe"] == pytest.approx(1.47)
        assert row["fitness"] == pytest.approx(1.20)
        assert row["turnover"] == pytest.approx(0.0383)
        assert row["long_count"] == 438
        assert row["remote_alpha_id"] == "kqVkNWmP"
        assert row["status"] == SimulationStatus.COMPLETED

    def test_failed_simulations_are_logged_too(self, tmp_path):
        path = tmp_path / "exp.xlsx"
        ExperimentLog(path).record(
            make_result(
                status=SimulationStatus.FAILED, grade=None, sharpe=None, fitness=None,
                passed=False, reasons=["Status FAILED != COMPLETED (parse error)"],
                error="parse error",
            ),
            SETTINGS,
        )

        _, rows = read_rows(path)
        assert len(rows) == 1
        assert rows[0]["status"] == SimulationStatus.FAILED
        assert rows[0]["error"] == "parse error"
        assert rows[0]["passed"] is False

    def test_settings_fall_back_to_the_result_when_not_passed(self, tmp_path):
        path = tmp_path / "exp.xlsx"
        ExperimentLog(path).record(make_result())

        _, rows = read_rows(path)
        assert rows[0]["neutralization"] == "SUBINDUSTRY"
        assert rows[0]["decay"] == 8

    def test_explicit_settings_win_over_the_result(self, tmp_path):
        path = tmp_path / "exp.xlsx"
        ExperimentLog(path).record(make_result(), {"region": "CHN", "decay": 3})

        _, rows = read_rows(path)
        assert rows[0]["region"] == "CHN"
        assert rows[0]["decay"] == 3

    def test_unwritable_path_does_not_raise(self, tmp_path):
        blocker = tmp_path / "blocker"
        blocker.write_text("i am a file", encoding="utf-8")
        ledger = ExperimentLog(blocker / "nested" / "exp.xlsx")
        assert ledger.record(make_result(), SETTINGS) is False

    def test_missing_openpyxl_would_not_break_a_run(self, tmp_path, monkeypatch):
        import builtins

        real_import = builtins.__import__

        def fake_import(name, *args, **kwargs):
            if name == "openpyxl":
                raise ImportError("no openpyxl")
            return real_import(name, *args, **kwargs)

        monkeypatch.setattr(builtins, "__import__", fake_import)
        ledger = ExperimentLog(tmp_path / "exp.xlsx")
        assert ledger.record(make_result(), SETTINGS) is False

    def test_outdated_header_migrates_values_by_name_not_position(self, tmp_path):
        # Regression: rewriting only the header left existing rows sitting under
        # the wrong column names. Inserting alpha_url after remote_alpha_id made
        # every older row show its region ("USA") in the alpha_url column.
        from openpyxl import Workbook

        path = tmp_path / "exp.xlsx"
        old_header = ["expression", "remote_alpha_id", "region", "sharpe", "dropped_col"]
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(old_header)
        sheet.append(["rank(close)", "OLD1", "USA", 1.41, "gone"])
        sheet.append(["rank(volume)", "OLD2", "CHN", 0.82, "gone"])
        workbook.save(path)

        assert ExperimentLog(path).record(make_result(), SETTINGS) is True

        header, rows = read_rows(path)
        assert header == list(LEDGER_COLUMNS)
        assert len(rows) == 3

        first, second, fresh = rows
        # Values stayed with their own columns despite the reordered old header.
        assert first["expression"] == "rank(close)"
        assert first["remote_alpha_id"] == "OLD1"
        assert first["region"] == "USA"
        assert first["sharpe"] == pytest.approx(1.41)
        assert second["region"] == "CHN"
        assert second["sharpe"] == pytest.approx(0.82)

        # The removed column is gone, and columns added since are blank for old rows.
        assert "dropped_col" not in header
        assert first["alpha_url"] in (None, "")
        assert first["test_sharpe"] in (None, "")

        # The newly appended row is complete and correctly aligned.
        assert fresh["remote_alpha_id"] == "kqVkNWmP"
        assert fresh["alpha_url"] == "https://platform.worldquantbrain.com/alpha/kqVkNWmP"
        assert fresh["region"] == "USA"
        assert fresh["sharpe"] == pytest.approx(1.47)

    def test_migration_of_an_empty_old_ledger_keeps_just_the_header(self, tmp_path):
        from openpyxl import Workbook

        path = tmp_path / "exp.xlsx"
        workbook = Workbook()
        workbook.active.append(["old_col_a", "old_col_b"])
        workbook.save(path)

        ExperimentLog(path).record(make_result(), SETTINGS)
        header, rows = read_rows(path)
        assert header == list(LEDGER_COLUMNS)
        assert len(rows) == 1

    def test_none_values_become_blank_cells(self, tmp_path):
        path = tmp_path / "exp.xlsx"
        ExperimentLog(path).record(make_result(sharpe=None, grade=None), SETTINGS)

        _, rows = read_rows(path)
        assert rows[0]["sharpe"] in (None, "")
        assert rows[0]["grade"] in (None, "")


class TestTestPeriodColumns:
    """A one-year test period is mandatory, so its results must reach the ledger."""

    TEST_STATS = {
        "sharpe": 0.42, "fitness": 0.13, "turnover": 0.0522, "returns": 0.0113,
        "drawdown": 0.083, "start_date": "2022-12-31",
    }
    TRAIN_STATS = {"sharpe": -0.53, "fitness": -0.18, "start_date": "2019-01-01"}

    def row_for(self, tmp_path, result):
        path = tmp_path / "exp.xlsx"
        assert ExperimentLog(path).record(result, SETTINGS) is True
        _, rows = read_rows(path)
        return rows[0]

    def test_alpha_url_is_built_from_the_remote_id(self, tmp_path):
        row = self.row_for(tmp_path, make_result())
        assert row["alpha_url"] == "https://platform.worldquantbrain.com/alpha/kqVkNWmP"

    def test_alpha_url_is_blank_without_a_remote_id(self, tmp_path):
        row = self.row_for(tmp_path, make_result(remote_alpha_id=None))
        assert row["alpha_url"] in (None, "")

    def test_train_and_test_metrics_are_recorded(self, tmp_path):
        row = self.row_for(tmp_path, make_result(
            train_stats=self.TRAIN_STATS, test_stats=self.TEST_STATS
        ))
        assert row["train_sharpe"] == pytest.approx(-0.53)
        assert row["train_fitness"] == pytest.approx(-0.18)
        assert row["test_sharpe"] == pytest.approx(0.42)
        assert row["test_fitness"] == pytest.approx(0.13)
        assert row["test_returns"] == pytest.approx(0.0113)
        assert row["test_turnover"] == pytest.approx(0.0522)
        assert row["test_drawdown"] == pytest.approx(0.083)

    def test_columns_stay_blank_without_a_test_period(self, tmp_path):
        row = self.row_for(tmp_path, make_result())
        for column in ("train_sharpe", "train_fitness", "test_sharpe", "test_fitness",
                       "test_returns", "test_turnover", "test_drawdown"):
            assert row[column] in (None, ""), f"{column} should be blank"

    def test_partially_populated_test_block(self, tmp_path):
        row = self.row_for(tmp_path, make_result(test_stats={"sharpe": 0.9}))
        assert row["test_sharpe"] == pytest.approx(0.9)
        assert row["test_fitness"] in (None, "")

    def test_diverging_test_period_is_visible_for_overfitting_review(self, tmp_path):
        # The whole point of holding out a year: a strong IS result with a weak
        # test year must be plainly visible side by side.
        row = self.row_for(tmp_path, make_result(
            sharpe=2.20, fitness=2.07, test_stats={"sharpe": -1.50, "fitness": -0.80}
        ))
        assert row["sharpe"] == pytest.approx(2.20)
        assert row["test_sharpe"] == pytest.approx(-1.50)


class TestSubmissionColumns:
    """Only an alpha with all eight checks PASS may be saved as a find."""

    def row_for(self, tmp_path, result):
        path = tmp_path / "exp.xlsx"
        assert ExperimentLog(path).record(result, SETTINGS) is True
        _, rows = read_rows(path)
        return rows[0]

    def test_the_columns_exist(self):
        assert {"submittable", "checks_passed", "self_correlation"} <= set(LEDGER_COLUMNS)

    def test_a_verified_alpha_reads_yes_and_eight_of_eight(self, tmp_path):
        row = self.row_for(tmp_path, make_result(
            submission_checks=all_pass_checks(), self_correlation=0.6996
        ))
        assert row["submittable"] == "YES"
        assert row["checks_passed"] == f"8/{len(SUBMISSION_CHECKS)}"
        assert row["self_correlation"] == pytest.approx(0.6996)
        assert row["checks_failed"] in (None, "")

    def test_a_failed_check_reads_no_and_names_the_reason(self, tmp_path):
        row = self.row_for(tmp_path, make_result(
            submission_checks=all_pass_checks({"SELF_CORRELATION": "FAIL"}),
            self_correlation=0.81,
        ))
        assert row["submittable"] == "NO"
        assert row["checks_passed"] == f"7/{len(SUBMISSION_CHECKS)}"
        assert "SELF_CORRELATION" in row["checks_failed"]

    def test_never_checked_leaves_the_verdict_blank(self, tmp_path):
        # Blank, not NO: a cell reading NO would claim the check ran and failed.
        row = self.row_for(tmp_path, make_result())
        assert row["submittable"] in (None, "")
        assert row["checks_passed"] in (None, "")
        assert row["self_correlation"] in (None, "")

    def test_a_pending_check_is_counted_as_not_passed(self, tmp_path):
        row = self.row_for(tmp_path, make_result(
            submission_checks=all_pass_checks({"SELF_CORRELATION": "PENDING"})
        ))
        assert row["submittable"] == "NO"
        assert row["checks_passed"] == f"7/{len(SUBMISSION_CHECKS)}"

    def test_an_unresolved_correlation_stays_blank(self, tmp_path):
        row = self.row_for(tmp_path, make_result(
            submission_checks=all_pass_checks(), self_correlation=None
        ))
        assert row["submittable"] == "YES"
        assert row["self_correlation"] in (None, "")


class TestBulkWrite:
    def test_writes_every_result(self, tmp_path):
        path = tmp_path / "bulk.xlsx"
        results = [make_result(alpha_id=f"a{i}") for i in range(4)]
        assert write_ledger_from_results(path, results) == 4

        _, rows = read_rows(path)
        assert len(rows) == 4

    def test_empty_input_creates_a_header_only_file(self, tmp_path):
        path = tmp_path / "empty.xlsx"
        assert write_ledger_from_results(path, []) == 0
        header, rows = read_rows(path)
        assert header == list(LEDGER_COLUMNS)
        assert rows == []


class TestBuildLedger:
    def test_wires_the_catalog_to_the_client(self, tmp_path):
        class Client:
            def __init__(self):
                self.calls = []

            def get_data_field(self, field_id):
                self.calls.append(field_id)
                return {"dataset": {"id": "model16"}}

        config = AppConfig(storage=StorageConfig(
            db_path=tmp_path / "t.db", data_dir=tmp_path,
            log_file=tmp_path / "t.log",
            ledger_path=tmp_path / "exp.xlsx",
            field_catalog_path=tmp_path / "cat.json",
        ))
        client = Client()
        ledger = build_ledger(config, client)
        assert ledger.record(make_result(), SETTINGS) is True
        assert client.calls == ["fn_liab_fair_val_l1_a"]

        _, rows = read_rows(config.storage.ledger_path)
        assert rows[0]["datasets"] == "model16"

    def test_tolerates_a_client_without_the_resolver(self, tmp_path):
        config = AppConfig(storage=StorageConfig(
            db_path=tmp_path / "t.db", data_dir=tmp_path, log_file=tmp_path / "t.log",
            ledger_path=tmp_path / "exp.xlsx", field_catalog_path=tmp_path / "cat.json",
        ))
        ledger = build_ledger(config, object())
        assert ledger.record(make_result(), SETTINGS) is True

        _, rows = read_rows(config.storage.ledger_path)
        assert rows[0]["fields"] == "fn_liab_fair_val_l1_a"
        assert rows[0]["datasets"] in (None, "")

    def test_works_without_any_client(self, tmp_path):
        config = AppConfig(storage=StorageConfig(
            db_path=tmp_path / "t.db", data_dir=tmp_path, log_file=tmp_path / "t.log",
            ledger_path=tmp_path / "exp.xlsx", field_catalog_path=tmp_path / "cat.json",
        ))
        assert build_ledger(config).record(make_result(), SETTINGS) is True
