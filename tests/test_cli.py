"""Command-line entry point tests.

The real ``scripts/run_backtest.py`` is imported and invoked; only the
``WorldQuantClient`` construction is patched so a fake BRAIN backend is injected.
"""

from __future__ import annotations

import csv
import sys
from pathlib import Path

import pytest

from conftest import FakeClock, FakeSession
from test_integration import EXPRESSIONS, BrainBackend
from worldquant.api import SimulationStatus
from worldquant.client import WorldQuantClient

PROJECT_ROOT = Path(__file__).resolve().parent.parent
SCRIPTS_DIR = PROJECT_ROOT / "scripts"
if str(SCRIPTS_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPTS_DIR))

import run_backtest  # noqa: E402


@pytest.fixture
def cli(tmp_path, monkeypatch):
    """Patch the client factory and credentials, and point all outputs at tmp_path."""
    clock = FakeClock()
    # Credential and .env auto-discovery resolve against config.PROJECT_ROOT, so
    # without this the tests would pick up the developer's real credentials.json
    # from the repository root. run_backtest's own PROJECT_ROOT (used only to
    # bootstrap sys.path) must stay untouched.
    monkeypatch.setattr("worldquant.config.PROJECT_ROOT", tmp_path)
    # The CLI builds its runner without an injectable clock, so a backend that
    # needs several polls would really sleep. Complete on the first poll here;
    # polling cadence is covered with a virtual clock in test_timeout.py.
    active = [BrainBackend(polls_before_done=0)]

    def install(backend):
        active[0] = backend

    def factory(credentials, **kwargs):
        kwargs.pop("min_request_interval", None)
        return WorldQuantClient(
            credentials,
            session=FakeSession(lambda method, url, kw: active[0](method, url, kw)),
            sleep=clock.sleep,
            clock=clock.now,
            min_request_interval=0.0,
            **kwargs,
        )

    monkeypatch.setattr(run_backtest, "WorldQuantClient", factory)
    monkeypatch.setenv("WQBRAIN_USERNAME", "tester@example.com")
    monkeypatch.setenv("WQBRAIN_PASSWORD", "s3cret")

    common = [
        "--db", str(tmp_path / "cli.db"),
        "--data-dir", str(tmp_path),
        "--log-file", str(tmp_path / "cli.log"),
        "--min-request-interval", "0",
        "--poll-interval", "1",
        "--poll-jitter", "0",
        "--max-wait", "300",
    ]

    def run(*args, env=None):
        for key, value in (env or {}).items():
            monkeypatch.setenv(key, value)
        return run_backtest.main([*args, *common])

    return {"run": run, "install": install, "backend": active[0], "clock": clock,
            "tmp_path": tmp_path, "monkeypatch": monkeypatch}


def read_csv(path):
    with Path(path).open(encoding="utf-8", newline="") as handle:
        return list(csv.DictReader(handle))


class TestExportOnly:
    def test_empty_database_exits_cleanly(self, cli, capsys):
        assert cli["run"]("--export-only") == 0
        assert (cli["tmp_path"] / "results.csv").exists()
        assert "No completed alphas to rank." in capsys.readouterr().out

    def test_needs_no_credentials(self, cli, monkeypatch):
        monkeypatch.delenv("WQBRAIN_USERNAME", raising=False)
        monkeypatch.delenv("WQBRAIN_PASSWORD", raising=False)
        assert cli["run"]("--export-only") == 0

    def test_creates_all_three_files(self, cli):
        cli["run"]("--export-only")
        for name in ("results.csv", "passed.csv", "failed.csv"):
            assert (cli["tmp_path"] / name).exists()


class TestRefreshYearly:
    """--refresh-yearly repairs rows whose yearly stats were silently dropped."""

    SETTINGS = {"region": "USA", "delay": 1}

    def seed(self, cli, *entries):
        """Store one completed alpha per ``(alpha_id, expression, grade)``."""
        from worldquant.hashing import dedup_key
        from worldquant.models import AlphaResult
        from worldquant.storage import ResultStore

        with ResultStore(cli["tmp_path"] / "cli.db") as store:
            for alpha_id, expression, grade in entries:
                alpha_row = store.upsert_alpha(expression, self.SETTINGS, name=alpha_id)
                sim_row = store.create_simulation(
                    alpha_row, remote_simulation_id=f"sim_{alpha_id}"
                )
                store.update_simulation(
                    sim_row, status=SimulationStatus.COMPLETED,
                    remote_alpha_id=alpha_id, mark_completed=True,
                )
                store.save_result(sim_row, AlphaResult(
                    alpha_id=alpha_id, expression=expression,
                    dedup_key=dedup_key(expression, self.SETTINGS),
                    status=SimulationStatus.COMPLETED, remote_alpha_id=alpha_id,
                    sharpe=1.4, grade=grade,
                ))

    def stored(self, cli):
        from worldquant.storage import ResultStore

        with ResultStore(cli["tmp_path"] / "cli.db") as store:
            return {r.remote_alpha_id: r for r in store.all_results()}

    def test_fills_the_missing_recordset(self, cli):
        backend = BrainBackend(polls_before_done=0)
        cli["install"](backend)
        self.seed(cli, ("A1", "rank(close)", "GOOD"))

        assert cli["run"]("--refresh-yearly") == 0
        assert self.stored(cli)["A1"].yearly_stats, "yearly stats should be backfilled"
        assert backend.submissions == 0, "no simulation may be submitted"

    def test_nothing_to_do_exits_cleanly(self, cli):
        backend = BrainBackend(polls_before_done=0)
        cli["install"](backend)
        assert cli["run"]("--refresh-yearly") == 0
        assert backend.submissions == 0

    def test_rows_that_already_have_stats_are_left_alone(self, cli):
        backend = BrainBackend(polls_before_done=0)
        cli["install"](backend)
        # A normal run populates yearly stats, so a refresh has nothing to do.
        cli["run"]("--expression", "rank(ts_delta(close, 5))")
        before = self.stored(cli)
        assert all(r.yearly_stats for r in before.values())

        assert cli["run"]("--refresh-yearly") == 0
        assert self.stored(cli) == before

    def test_limit_caps_the_requests_and_prefers_the_best_grades(self, cli):
        backend = BrainBackend(polls_before_done=0)
        cli["install"](backend)
        self.seed(cli, ("LOW1", "rank(low)", "INFERIOR"), ("HI1", "rank(hi)", "SPECTACULAR"))

        assert cli["run"]("--refresh-yearly", "--limit", "1") == 0
        stored = self.stored(cli)
        assert stored["HI1"].yearly_stats, "the better-graded alpha goes first"
        assert not stored["LOW1"].yearly_stats

    def test_an_alpha_without_a_remote_id_is_skipped(self, cli):
        from worldquant.hashing import dedup_key
        from worldquant.models import AlphaResult
        from worldquant.storage import ResultStore

        backend = BrainBackend(polls_before_done=0)
        cli["install"](backend)
        with ResultStore(cli["tmp_path"] / "cli.db") as store:
            alpha_row = store.upsert_alpha("rank(close)", self.SETTINGS, name="orphan")
            sim_row = store.create_simulation(alpha_row, remote_simulation_id="sim_x")
            store.update_simulation(sim_row, status=SimulationStatus.COMPLETED)
            store.save_result(sim_row, AlphaResult(
                alpha_id="orphan", expression="rank(close)",
                dedup_key=dedup_key("rank(close)", self.SETTINGS),
                status=SimulationStatus.COMPLETED, remote_alpha_id=None, grade="GOOD",
            ))

        assert cli["run"]("--refresh-yearly") == 0
        assert backend.submissions == 0

    def test_an_unavailable_endpoint_is_reported_not_fatal(self, cli):
        backend = BrainBackend(polls_before_done=0, yearly_status=404, yearly_payload={})
        cli["install"](backend)
        self.seed(cli, ("A1", "rank(close)", "GOOD"))

        assert cli["run"]("--refresh-yearly") == 0
        assert not self.stored(cli)["A1"].yearly_stats


class TestSingleExpression:
    def test_runs_one_alpha(self, cli, capsys):
        code = cli["run"]("--expression", "rank(ts_delta(close, 5))")
        assert code == 0
        assert cli["backend"].submissions == 1

        rows = read_csv(cli["tmp_path"] / "results.csv")
        assert len(rows) == 1
        assert rows[0]["expression"] == "rank(ts_delta(close, 5))"
        assert rows[0]["status"] == SimulationStatus.COMPLETED
        assert float(rows[0]["sharpe"]) == pytest.approx(1.41)

    def test_expression_is_repeatable(self, cli):
        code = cli["run"]("-e", "rank(close)", "-e", "rank(volume)")
        assert code == 0
        assert cli["backend"].submissions == 2

    def test_leaderboard_is_printed(self, cli, capsys):
        cli["run"]("--expression", "rank(close)")
        out = capsys.readouterr().out
        assert "Top 1 Alpha" in out
        assert "Sharpe" in out


class TestBatchInput:
    def test_csv_input(self, cli, tmp_path):
        source = tmp_path / "alphas.csv"
        source.write_text(
            'name,expression\na1,"rank(ts_delta(close, 5))"\na2,"rank(ts_mean(volume, 20))"\n',
            encoding="utf-8",
        )
        assert cli["run"]("--input", str(source)) == 0
        assert cli["backend"].submissions == 2
        assert len(read_csv(tmp_path / "results.csv")) == 2

    def test_txt_input(self, cli, tmp_path):
        source = tmp_path / "alphas.txt"
        source.write_text("rank(close)\nrank(volume)\n", encoding="utf-8")
        assert cli["run"]("--input", str(source)) == 0
        assert cli["backend"].submissions == 2

    def test_limit_caps_the_batch(self, cli, tmp_path):
        source = tmp_path / "alphas.txt"
        source.write_text("\n".join(EXPRESSIONS) + "\n", encoding="utf-8")
        assert cli["run"]("--input", str(source), "--limit", "2") == 0
        assert cli["backend"].submissions == 2

    def test_expression_and_input_are_merged_and_deduplicated(self, cli, tmp_path):
        source = tmp_path / "alphas.txt"
        source.write_text("rank(close)\n", encoding="utf-8")
        assert cli["run"]("--input", str(source), "-e", "rank( close )") == 0
        assert cli["backend"].submissions == 1

    def test_missing_input_file_is_a_usage_error(self, cli, tmp_path):
        assert cli["run"]("--input", str(tmp_path / "nope.csv")) == 2


class TestResumeAndForce:
    def test_second_run_submits_nothing(self, cli, tmp_path):
        source = tmp_path / "alphas.txt"
        source.write_text("\n".join(EXPRESSIONS) + "\n", encoding="utf-8")
        assert cli["run"]("--input", str(source)) == 0
        assert cli["backend"].submissions == 3

        assert cli["run"]("--input", str(source)) == 0
        assert cli["backend"].submissions == 3

    def test_force_resubmits(self, cli, tmp_path):
        source = tmp_path / "alphas.txt"
        source.write_text("rank(close)\n", encoding="utf-8")
        cli["run"]("--input", str(source))
        assert cli["run"]("--input", str(source), "--force") == 0
        assert cli["backend"].submissions == 2

    def test_resume_only_finishes_in_flight_work(self, cli):
        from worldquant.storage import ResultStore

        # Leave a RUNNING simulation behind, as an interrupted run would.
        with ResultStore(cli["tmp_path"] / "cli.db") as store:
            alpha_row = store.upsert_alpha("rank(close)", {"region": "USA", "delay": 1}, name="a1")
            store.create_simulation(
                alpha_row, remote_simulation_id="sim999", status=SimulationStatus.RUNNING
            )

        assert cli["run"]("--resume-only") == 0
        assert cli["backend"].submissions == 0
        assert "sim999" in cli["backend"].poll_counts

        rows = read_csv(cli["tmp_path"] / "results.csv")
        assert len(rows) == 1
        assert rows[0]["status"] == SimulationStatus.COMPLETED

    def test_resume_only_with_nothing_to_do(self, cli):
        assert cli["run"]("--resume-only") == 0


class TestCredentialsFlag:
    def write_creds(self, tmp_path, email="file@example.com", password="file-secret",
                    name="creds.json"):
        import json

        path = tmp_path / name
        path.write_text(json.dumps({"email": email, "password": password}), encoding="utf-8")
        return path

    def test_supplies_credentials_when_no_env_vars_are_set(self, cli, monkeypatch):
        monkeypatch.delenv("WQBRAIN_USERNAME", raising=False)
        monkeypatch.delenv("WQBRAIN_PASSWORD", raising=False)
        path = self.write_creds(cli["tmp_path"])

        assert cli["run"]("--expression", "rank(close)", "--credentials", str(path)) == 0
        assert cli["backend"].submissions == 1

    def test_takes_precedence_over_env_vars(self, cli, monkeypatch):
        monkeypatch.setenv("WQBRAIN_USERNAME", "env@example.com")
        monkeypatch.setenv("WQBRAIN_PASSWORD", "env-secret")
        path = self.write_creds(cli["tmp_path"])

        assert cli["run"]("--expression", "rank(close)", "--credentials", str(path)) == 0
        assert cli["backend"].submissions == 1

    def test_missing_file_is_a_usage_error(self, cli, tmp_path):
        code = cli["run"]("--expression", "rank(close)",
                          "--credentials", str(tmp_path / "gone.json"))
        assert code == 2
        assert cli["backend"].submissions == 0

    def test_unfilled_template_is_a_usage_error(self, cli, tmp_path):
        import json

        path = tmp_path / "template.json"
        path.write_text(
            json.dumps({"email": "you@example.com", "password": "your-password"}),
            encoding="utf-8",
        )
        assert cli["run"]("--expression", "rank(close)", "--credentials", str(path)) == 2
        assert cli["backend"].submissions == 0

    def test_malformed_json_is_a_usage_error(self, cli, tmp_path):
        path = tmp_path / "broken.json"
        path.write_text('{"email": "me@example.com", "password": ', encoding="utf-8")
        assert cli["run"]("--expression", "rank(close)", "--credentials", str(path)) == 2
        assert cli["backend"].submissions == 0

    def test_password_from_the_file_is_never_logged(self, cli):
        path = self.write_creds(cli["tmp_path"], password="sup3r-s3kr3t")
        cli["run"]("--expression", "rank(close)", "--credentials", str(path))
        log_text = (cli["tmp_path"] / "cli.log").read_text(encoding="utf-8")
        assert "sup3r-s3kr3t" not in log_text
        assert "file@example.com" in log_text


class TestUsageErrors:
    def test_no_alpha_source_is_a_usage_error(self, cli):
        assert cli["run"]() == 2

    def test_missing_credentials_is_a_usage_error(self, cli, monkeypatch):
        monkeypatch.delenv("WQBRAIN_USERNAME", raising=False)
        monkeypatch.delenv("WQBRAIN_PASSWORD", raising=False)
        assert cli["run"]("--expression", "rank(close)") == 2

    def test_half_configured_credentials_is_a_usage_error(self, cli, monkeypatch):
        monkeypatch.delenv("WQBRAIN_PASSWORD", raising=False)
        assert cli["run"]("--expression", "rank(close)") == 2

    def test_bogus_config_file_is_a_usage_error(self, cli, tmp_path):
        bad = tmp_path / "bad.yaml"
        bad.write_text("filters:\n  min_sharpe_ratio: 1\n", encoding="utf-8")
        assert cli["run"]("--expression", "rank(close)", "--config", str(bad)) == 2

    def test_missing_config_file_is_a_usage_error(self, cli, tmp_path):
        assert cli["run"]("--expression", "rank(close)", "--config", str(tmp_path / "no.yaml")) == 2


class TestRuntimeFailureExitCode:
    def test_auth_failure_exits_nonzero(self, cli, monkeypatch):
        from worldquant.exceptions import AuthError

        def failing_factory(credentials, **kwargs):
            class Failing:
                def __enter__(self):
                    return self

                def __exit__(self, *exc):
                    return False

                def ensure_authenticated(self):
                    raise AuthError("bad credentials", status_code=401, url="/authentication")

            return Failing()

        monkeypatch.setattr(run_backtest, "WorldQuantClient", failing_factory)
        assert cli["run"]("--expression", "rank(close)") == 1


class TestConcurrencyClamping:
    def test_excessive_concurrency_is_clamped_not_rejected(self, cli, tmp_path):
        source = tmp_path / "alphas.txt"
        source.write_text("\n".join(EXPRESSIONS) + "\n", encoding="utf-8")
        assert cli["run"]("--input", str(source), "--concurrency", "50") == 0
        assert cli["backend"].submissions == 3


class TestExperimentLedger:
    def ledger_path(self, cli):
        return cli["tmp_path"] / "data" / "experiments.xlsx"

    def rows(self, cli):
        from openpyxl import load_workbook

        sheet = load_workbook(self.ledger_path(cli)).active
        header = [cell.value for cell in sheet[1]]
        return [dict(zip(header, raw)) for raw in sheet.iter_rows(min_row=2, values_only=True)]

    def test_a_run_records_one_row_per_alpha(self, cli, tmp_path):
        source = tmp_path / "alphas.txt"
        source.write_text("rank(close)\nrank(volume)\n", encoding="utf-8")
        assert cli["run"]("--input", str(source)) == 0

        rows = self.rows(cli)
        assert len(rows) == 2
        assert {row["expression"] for row in rows} == {"rank(close)", "rank(volume)"}

    def test_row_carries_fields_and_parameters(self, cli):
        cli["run"]("--expression", "rank(ts_delta(close, 5))")
        row = self.rows(cli)[0]
        assert row["fields"] == "close"
        assert row["field_count"] == 1
        assert row["region"] == "USA"
        assert row["delay"] == 1
        assert row["neutralization"] == "INDUSTRY"
        assert row["language"] == "FASTEXPR"

    def test_row_carries_results_and_grade(self, cli):
        cli["run"]("--expression", "rank(close)")
        row = self.rows(cli)[0]
        assert row["grade"] == "INFERIOR"
        assert row["status"] == SimulationStatus.COMPLETED
        assert float(row["sharpe"]) == pytest.approx(1.41)
        assert row["remote_alpha_id"]

    def test_failed_simulations_are_recorded_too(self, cli):
        from conftest import FakeResponse

        inner = BrainBackend(polls_before_done=0)

        class FailingBackend:
            """Rejects every submission, as BRAIN does for an invalid expression."""

            def __call__(self, method, url, kwargs):
                if url.endswith("/simulations") and method == "POST":
                    return FakeResponse(400, json_data={"detail": "unknown operator"})
                return inner(method, url, kwargs)

        cli["install"](FailingBackend())
        cli["run"]("--expression", "bogus(close)")

        rows = self.rows(cli)
        assert len(rows) == 1
        assert rows[0]["status"] == SimulationStatus.REQUEST_ERROR
        assert "unknown operator" in rows[0]["error"]
        # A failure still records what was attempted.
        assert rows[0]["expression"] == "bogus(close)"
        assert rows[0]["fields"] == "close"

    def test_rebuild_ledger_regenerates_from_the_database(self, cli):
        cli["run"]("--expression", "rank(close)")
        path = self.ledger_path(cli)
        path.unlink()
        assert not path.exists()

        assert cli["run"]("--export-only", "--rebuild-ledger") == 0
        rows = self.rows(cli)
        assert len(rows) == 1
        assert rows[0]["expression"] == "rank(close)"

    def test_rebuilding_twice_does_not_duplicate_rows(self, cli):
        cli["run"]("--expression", "rank(close)")
        cli["run"]("--export-only", "--rebuild-ledger")
        cli["run"]("--export-only", "--rebuild-ledger")
        assert len(self.rows(cli)) == 1

    def test_export_only_without_the_flag_leaves_the_ledger_alone(self, cli):
        cli["run"]("--expression", "rank(close)")
        path = self.ledger_path(cli)
        before = path.read_bytes()
        cli["run"]("--export-only")
        assert path.read_bytes() == before


class TestLogFile:
    def test_log_file_is_written(self, cli):
        cli["run"]("--expression", "rank(close)")
        log_path = cli["tmp_path"] / "cli.log"
        assert log_path.exists()
        content = log_path.read_text(encoding="utf-8")
        assert "[INFO]" in content
        assert "Login successful" in content

    def test_password_is_never_logged(self, cli):
        cli["run"]("--expression", "rank(close)")
        assert "s3cret" not in (cli["tmp_path"] / "cli.log").read_text(encoding="utf-8")
