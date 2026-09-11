"""Experiment ledger: one xlsx row per simulation.

Records what was actually tried — dataset(s), data field(s), every backtest
parameter, and the full result — so a search can be reviewed, compared and
reproduced without digging through SQLite or the platform UI.

The ledger is append-only and network-free. Dataset names come from
:class:`FieldCatalog`, which caches a field -> dataset mapping in a local JSON
file and can be refreshed from the API when a client is supplied.
"""

from __future__ import annotations

import json
import re
import threading
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Callable, Iterable, Sequence

from .api import SUBMISSION_CHECKS, passed_check_count
from .exceptions import StorageError
from .logging_utils import get_logger
from .models import AlphaResult

#: Words that are operators/functions/keywords rather than data fields. Anything
#: immediately followed by ``(`` is excluded by :func:`extract_field_names` too.
_NON_FIELD_WORDS = frozenset(
    {
        "if", "else", "and", "or", "not", "true", "false", "null", "nan",
        "group_neutralize", "group_rank", "group_zscore", "rank", "zscore",
        "scale", "scale_down", "signed_power", "power", "abs", "log", "sign",
        "max", "min", "sum", "avg", "stddev", "corr", "covariance", "delta",
        "decay_linear", "ts_rank", "ts_mean", "ts_sum", "ts_std_dev", "ts_delta",
        "ts_max", "ts_min", "ts_arg_max", "ts_arg_min", "ts_corr", "ts_covariance",
        "ts_av_diff", "ts_regression", "ts_zscore", "ts_scale", "ts_product",
        "ts_kurtosis", "ts_skewness", "ts_decay_linear", "ts_backfill",
        "bucket", "keep", "when", "indclass", "sector", "industry", "subindustry",
        "market", "country", "exchange", "currency", "vega", "beta",
    }
)

_IDENTIFIER = re.compile(r"[A-Za-z_][A-Za-z0-9_]*")

#: Platform link prefix for an alpha, used by the ledger's ``alpha_url`` column.
ALPHA_URL_PREFIX = "https://platform.worldquantbrain.com/alpha/"

#: Ledger columns, in order. One row per simulation.
LEDGER_COLUMNS: tuple[str, ...] = (
    "run_at",
    "alpha_id",
    "status",
    "grade",
    "expression",
    "datasets",
    "fields",
    "field_count",
    "simulation_id",
    "remote_alpha_id",
    "alpha_url",
    "region",
    "universe",
    "delay",
    "decay",
    "neutralization",
    "truncation",
    "pasteurization",
    "nan_handling",
    "unit_handling",
    "instrument_type",
    "language",
    "sharpe",
    "fitness",
    "turnover",
    "returns",
    "drawdown",
    "margin",
    "pnl",
    "book_size",
    "long_count",
    "short_count",
    "train_sharpe",
    "train_fitness",
    "test_sharpe",
    "test_fitness",
    "test_returns",
    "test_turnover",
    "test_drawdown",
    "checks_failed",
    "submittable",
    "checks_passed",
    "self_correlation",
    "excluded_reason",
    "passed",
    "reasons",
    "positive_years",
    "negative_years",
    "total_years",
    "positive_year_ratio",
    "yearly_sharpe_std",
    "worst_year_sharpe",
    "error",
)

#: settings key -> ledger column
_SETTINGS_COLUMNS: dict[str, str] = {
    "region": "region",
    "universe": "universe",
    "delay": "delay",
    "decay": "decay",
    "neutralization": "neutralization",
    "truncation": "truncation",
    "pasteurization": "pasteurization",
    "nanHandling": "nan_handling",
    "unitHandling": "unit_handling",
    "instrumentType": "instrument_type",
    "language": "language",
}


def extract_field_names(expression: str) -> list[str]:
    """Pull data-field identifiers out of an expression.

    A token counts as a field when it is not a function name (immediately
    followed by ``(``), not a named argument (``std=3.0``), not inside a string
    literal, and not a local variable of a multi-statement expression.

    Multi-statement FASTEXPR binds intermediates::

        ey = ts_backfill(anl4_ebit_value, 40) / cap;
        raw = group_zscore(ey, industry)

    ``ey`` and ``raw`` are variables, not dataset fields; reporting them would
    pollute the ledger's provenance columns and waste ``/data-fields`` lookups.
    Only an identifier at parenthesis depth 0 followed by a bare ``=`` is a
    definition — named arguments live inside a call (depth >= 1) and are excluded
    by a separate rule. Order of first appearance is kept, duplicates collapse.
    """
    if not expression:
        return []

    matches = list(_IDENTIFIER.finditer(expression))
    in_string = _string_mask(expression, matches)
    depths = _paren_depths(expression, matches)

    assigned: set[str] = set()
    for index, match in enumerate(matches):
        if in_string[index] or depths[index] != 0:
            continue
        if _is_assignment_target(expression, match.end()):
            assigned.add(match.group(0).lower())

    fields: list[str] = []
    seen: set[str] = set()
    for index, match in enumerate(matches):
        if in_string[index]:
            continue
        token = match.group(0)
        lowered = token.lower()
        if lowered in _NON_FIELD_WORDS or lowered in seen or lowered in assigned:
            continue
        # Function call.
        if expression[match.end():].lstrip().startswith("("):
            continue
        # Named argument such as std=3.0 or RETTYPE=1.
        if _is_assignment_target(expression, match.end()):
            continue
        seen.add(lowered)
        fields.append(token)
    return fields


def _string_mask(expression: str, matches: list[re.Match]) -> list[bool]:
    """For each match, whether it starts inside a single- or double-quoted string."""
    return [
        expression[: match.start()].count("'") % 2 == 1
        or expression[: match.start()].count('"') % 2 == 1
        for match in matches
    ]


def _paren_depths(expression: str, matches: list[re.Match]) -> list[int]:
    """Parenthesis nesting depth at the start of each match, ignoring literals."""
    depths: list[int] = []
    for match in matches:
        depth = 0
        quote: str | None = None
        for char in expression[: match.start()]:
            if quote is not None:
                if char == quote:
                    quote = None
                continue
            if char in "'\"":
                quote = char
            elif char == "(":
                depth += 1
            elif char == ")":
                depth = max(0, depth - 1)
        depths.append(depth)
    return depths


def _is_assignment_target(expression: str, end: int) -> bool:
    """True when the identifier ending at ``end`` is followed by a bare ``=``.

    ``==`` is a comparison, so the operand before it is still a field.
    """
    stripped = expression[end:].lstrip()
    return stripped.startswith("=") and not stripped.startswith("==")


class FieldCatalog:
    """Maps a data field id to its dataset id, cached in a local JSON file.

    Resolution needs the ``/data-fields/{id}`` endpoint, so a client accessor is
    injected rather than imported: the ledger itself never touches the network.
    Unknown or unresolvable fields simply record no dataset.
    """

    def __init__(
        self,
        cache_path: str | Path,
        *,
        fetch: Callable[[str], dict[str, Any] | None] | None = None,
        logger: Any = None,
    ) -> None:
        self.cache_path = Path(cache_path)
        self._fetch = fetch
        self.log = logger or get_logger("catalog")
        self._lock = threading.RLock()
        self._map: dict[str, str] = {}
        self._misses: set[str] = set()
        self._load()

    def _load(self) -> None:
        if not self.cache_path.exists():
            return
        try:
            payload = json.loads(self.cache_path.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError) as exc:
            self.log.warning("ignoring unreadable field catalog %s: %s", self.cache_path, exc)
            return
        if isinstance(payload, dict):
            self._map = {str(k): str(v) for k, v in payload.items() if isinstance(v, str)}

    def _save(self) -> None:
        try:
            self.cache_path.parent.mkdir(parents=True, exist_ok=True)
            self.cache_path.write_text(
                json.dumps(self._map, sort_keys=True, indent=1), encoding="utf-8"
            )
        except OSError as exc:
            self.log.warning("cannot persist field catalog %s: %s", self.cache_path, exc)

    def dataset_for(self, field: str) -> str | None:
        """Dataset id for a field, resolving through the API when uncached."""
        if not field:
            return None
        with self._lock:
            if field in self._map:
                return self._map[field]
            if field in self._misses or self._fetch is None:
                return None
            try:
                payload = self._fetch(field)
            except Exception as exc:  # noqa: BLE001 - isolation boundary
                # Dataset attribution is optional enrichment. Whatever the
                # resolver does — network error, unexpected exception, a stub
                # that refuses the URL — it must cost at most the dataset
                # column, never the whole ledger row.
                self.log.debug("cannot resolve dataset for %s: %s", field, exc)
                self._misses.add(field)
                return None

            dataset = None
            if isinstance(payload, dict):
                candidate = payload.get("dataset")
                if isinstance(candidate, dict):
                    dataset = candidate.get("id")
                elif isinstance(candidate, str):
                    dataset = candidate
            if isinstance(dataset, str) and dataset:
                self._map[field] = dataset
                self._save()
                return dataset
            self._misses.add(field)
            return None

    def datasets_for(self, fields: Iterable[str]) -> list[str]:
        """Distinct dataset ids for a set of fields, in first-seen order."""
        found: list[str] = []
        seen: set[str] = set()
        for field in fields:
            dataset = self.dataset_for(field)
            if dataset and dataset not in seen:
                seen.add(dataset)
                found.append(dataset)
        return found

    @property
    def size(self) -> int:
        return len(self._map)


def _utcnow() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


class ExperimentLog:
    """Append-only xlsx ledger, one row per simulation.

    Safe to share across the runner's worker threads. Writes are serialized and
    failures are logged rather than raised: losing a ledger row must never abort
    a backtest that already succeeded.
    """

    def __init__(
        self,
        path: str | Path,
        *,
        catalog: FieldCatalog | None = None,
        logger: Any = None,
    ) -> None:
        self.path = Path(path)
        self.catalog = catalog
        self.log = logger or get_logger("ledger")
        self._lock = threading.RLock()
        self._rows_written = 0

    def record(self, result: AlphaResult, settings: dict[str, Any] | None = None) -> bool:
        """Append one simulation to the ledger. Returns True on success."""
        try:
            row = self.build_row(result, settings)
        except Exception as exc:  # noqa: BLE001 - the ledger must never break a run
            self.log.warning("cannot build ledger row for %s: %s", result.alpha_id, exc)
            return False

        try:
            self._append(row)
        except Exception as exc:  # noqa: BLE001 - same reason
            self.log.warning("cannot write ledger %s: %s", self.path, exc)
            return False

        self._rows_written += 1
        return True

    def build_row(
        self, result: AlphaResult, settings: dict[str, Any] | None = None
    ) -> dict[str, Any]:
        """Flatten one result into ledger columns."""
        fields = extract_field_names(result.expression)
        datasets = self.catalog.datasets_for(fields) if self.catalog else []

        effective = dict(settings or {})
        if not effective and result.settings_json:
            try:
                parsed = json.loads(result.settings_json)
                if isinstance(parsed, dict):
                    effective = parsed
            except (json.JSONDecodeError, TypeError):
                effective = {}

        summary = result.yearly_summary
        row: dict[str, Any] = {column: "" for column in LEDGER_COLUMNS}
        row.update(
            {
                "run_at": result.completed_at or result.created_at or _utcnow(),
                "alpha_id": result.alpha_id,
                "status": result.status,
                "grade": result.grade or "",
                "expression": result.expression,
                "datasets": ", ".join(datasets),
                "fields": ", ".join(fields),
                "field_count": len(fields),
                "simulation_id": result.simulation_id or "",
                "remote_alpha_id": result.remote_alpha_id or "",
                "alpha_url": (
                    f"{ALPHA_URL_PREFIX}{result.remote_alpha_id}"
                    if result.remote_alpha_id else ""
                ),
                "sharpe": result.sharpe,
                "fitness": result.fitness,
                "turnover": result.turnover,
                "returns": result.returns,
                "drawdown": result.drawdown,
                "margin": result.margin,
                "pnl": result.pnl,
                "book_size": result.book_size,
                "long_count": result.long_count,
                "short_count": result.short_count,
                "train_sharpe": _stage_metric(result.train_stats, "sharpe"),
                "train_fitness": _stage_metric(result.train_stats, "fitness"),
                "test_sharpe": _stage_metric(result.test_stats, "sharpe"),
                "test_fitness": _stage_metric(result.test_stats, "fitness"),
                "test_returns": _stage_metric(result.test_stats, "returns"),
                "test_turnover": _stage_metric(result.test_stats, "turnover"),
                "test_drawdown": _stage_metric(result.test_stats, "drawdown"),
                "checks_failed": ", ".join(failed_checks(result)),
                "submittable": (
                    "" if not result.submission_checks
                    else ("YES" if result.is_submittable else "NO")
                ),
                "checks_passed": (
                    "" if not result.submission_checks
                    else f"{passed_check_count(result.submission_checks)}/{len(SUBMISSION_CHECKS)}"
                ),
                "self_correlation": (
                    "" if result.self_correlation is None else result.self_correlation
                ),
                "excluded_reason": result.excluded_reason or "",
                "passed": "" if result.passed is None else bool(result.passed),
                "reasons": "; ".join(result.reasons),
                "positive_years": summary.positive_years,
                "negative_years": summary.negative_years,
                "total_years": summary.total_years,
                "positive_year_ratio": summary.positive_year_ratio,
                "yearly_sharpe_std": summary.yearly_sharpe_std,
                "worst_year_sharpe": summary.worst_year_sharpe,
                "error": result.error or "",
            }
        )
        for settings_key, column in _SETTINGS_COLUMNS.items():
            if settings_key in effective:
                row[column] = effective[settings_key]
        # Blank out None so openpyxl writes an empty cell rather than "None".
        return {key: ("" if value is None else value) for key, value in row.items()}

    def ensure_file(self) -> Path:
        """Create the ledger with its header row if it does not exist yet.

        Guarantees a usable file even when there is nothing to record, so a
        backfill over an empty result set still leaves behind a valid ledger.
        """
        with self._lock:
            if not self.path.exists():
                self._append({column: "" for column in LEDGER_COLUMNS})
                # _append wrote a blank data row along with the header; drop it.
                from openpyxl import load_workbook

                workbook = load_workbook(self.path)
                sheet = workbook.active
                if sheet.max_row > 1:
                    sheet.delete_rows(2, sheet.max_row - 1)
                workbook.save(self.path)
        return self.path

    def _append(self, row: dict[str, Any]) -> None:
        try:
            from openpyxl import Workbook, load_workbook
        except ImportError as exc:
            raise StorageError(
                "openpyxl is required for the xlsx experiment ledger; "
                "pip install openpyxl"
            ) from exc

        with self._lock:
            self.path.parent.mkdir(parents=True, exist_ok=True)
            if self.path.exists():
                workbook = load_workbook(self.path)
                sheet = workbook.active
                header = [cell.value for cell in sheet[1]] if sheet.max_row >= 1 else []
                if header != list(LEDGER_COLUMNS):
                    # Rewriting only the header would leave every existing row
                    # sitting under the wrong column names: inserting a column
                    # mid-schema shifts all later values by one. Migrate by name
                    # instead, so each value stays with its field.
                    migrated = [
                        dict(zip(header, raw))
                        for raw in sheet.iter_rows(min_row=2, values_only=True)
                    ]
                    self.log.warning(
                        "ledger %s has an outdated header (%d column(s)); migrating "
                        "%d existing row(s) to the current %d-column layout",
                        self.path, len(header), len(migrated), len(LEDGER_COLUMNS),
                    )
                    sheet.delete_rows(1, sheet.max_row)
                    sheet.append(list(LEDGER_COLUMNS))
                    for old_row in migrated:
                        sheet.append([old_row.get(column, "") for column in LEDGER_COLUMNS])
            else:
                workbook = Workbook()
                sheet = workbook.active
                sheet.title = "simulations"
                sheet.append(list(LEDGER_COLUMNS))

            sheet.append([row.get(column, "") for column in LEDGER_COLUMNS])
            workbook.save(self.path)

    @property
    def rows_written(self) -> int:
        return self._rows_written


def _stage_metric(stage: dict[str, Any] | None, metric: str) -> Any:
    """One metric from a parsed train/test block, or "" when unavailable.

    Runs without a ``testPeriod`` have no train/test blocks at all, and those
    ledger cells should stay blank rather than read as a measured zero.
    """
    if not isinstance(stage, dict):
        return ""
    value = stage.get(metric)
    return "" if value is None else value


def failed_checks(result: AlphaResult) -> list[str]:
    """Names of the BRAIN IS checks that did not pass.

    These gate whether an alpha can be **submitted** to WorldQuant; they do not
    decide the grade (fitness does — see :class:`~worldquant.api.AlphaGrade`).
    Still the most useful diagnostic in the ledger, because a high grade with
    failing checks means the alpha cannot actually be submitted:
    ``LOW_FITNESS=0.83(limit 1.0):FAIL`` says exactly what is blocking it.

    Reads the checks resolved by ``GET /alphas/{id}/check`` when available. The
    alpha payload's own copy reports ``SELF_CORRELATION`` as ``PENDING`` forever,
    so a correlation failure would otherwise be invisible here while
    ``submittable`` said NO — a verdict with no reason attached.
    """
    failed: list[str] = []
    source = result.submission_checks or result.checks or {}
    for name, check in source.items():
        if not isinstance(check, dict):
            continue
        outcome = check.get("result")
        if isinstance(outcome, str) and outcome.upper() not in {"PASS", "PENDING"}:
            value = check.get("value")
            limit = check.get("limit")
            detail = f"{name}={value}(limit {limit})" if value is not None else name
            failed.append(f"{detail}:{outcome}")
    return failed


def build_ledger(config: Any, client: Any = None, *, logger: Any = None) -> ExperimentLog:
    """Build the ledger (and its dataset catalog) from an :class:`AppConfig`.

    When a client is supplied, uncached fields are resolved to their dataset via
    ``GET /data-fields/{id}`` and cached locally, so the lookup cost is paid once
    per field rather than once per simulation.
    """
    resolver = getattr(client, "get_data_field", None) if client is not None else None
    catalog = FieldCatalog(
        config.storage.field_catalog_path,
        fetch=resolver if callable(resolver) else None,
        logger=logger,
    )
    return ExperimentLog(config.storage.ledger_path, catalog=catalog, logger=logger)


def write_ledger_from_results(
    path: str | Path,
    results: Sequence[AlphaResult],
    *,
    catalog: FieldCatalog | None = None,
    logger: Any = None,
) -> int:
    """Bulk-write a ledger, e.g. to backfill runs recorded before it existed."""
    ledger = ExperimentLog(path, catalog=catalog, logger=logger)
    ledger.ensure_file()
    written = 0
    for result in results:
        if ledger.record(result):
            written += 1
    return written
